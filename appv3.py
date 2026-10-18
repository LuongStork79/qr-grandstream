from flask import Flask, render_template_string, request, jsonify
import openpyxl
import os
import sys
from datetime import datetime
import socket
import threading
import json
import base64
import re

app = Flask(__name__)

excel_lock = threading.Lock()

if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

EXCEL_FILE = os.path.join(BASE_DIR, "DanhSachThietBi.xlsx")
pending_queue = []
recent_duplicates = []

# Tự động tìm chính xác file LogoCLG .png trong cùng thư mục
LOGO_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
for fname in ["LogoCLG .png", "LogoCLG.png", "Logo CLG (new).png", "Logo CLG (new) (1).jpg", "logo.png", "logo.jpg"]:
    logo_path = os.path.join(BASE_DIR, fname)
    if os.path.exists(logo_path):
        try:
            with open(logo_path, "rb") as f:
                LOGO_BASE64 = base64.b64encode(f.read()).decode('utf-8')
                print(f"[*] Đã tải thành công logo từ file: {fname}")
                break
        except Exception as e:
            print(f"[!] Lỗi đọc file logo {fname}: {e}")

def init_excel_file():
    with excel_lock:
        if not os.path.exists(EXCEL_FILE):
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Data Thiết Bị"
                headers = ["STT", "MAC Address", "Serial Number", "Password", "Thời Gian"]
                ws.append(headers)
                ws.column_dimensions['A'].width = 10
                ws.column_dimensions['B'].width = 25
                ws.column_dimensions['C'].width = 25
                ws.column_dimensions['D'].width = 20
                ws.column_dimensions['E'].width = 20
                wb.save(EXCEL_FILE)
                print(f"[*] Đã tự động sinh mới file Excel tại: {EXCEL_FILE}")
            except Exception as e:
                print(f"[!] Không thể tạo file Excel: {e}")

def get_total_devices():
    init_excel_file()
    with excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True)
            ws = wb.active
            total = max(0, ws.max_row - 1)
            return total + len(pending_queue)
        except Exception:
            return len(pending_queue)

VIEW_TEMPLATE = """
<!DOCTYPE html>
<html lang="vi" data-theme="light">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Qr Grandstream by CLG - Enterprise Dashboard</title>
    <!-- Google Fonts: Inter & JetBrains Mono -->
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=JetBrains+Mono:wght@600;700&display=swap" rel="stylesheet">
    <style>
        /* Định nghĩa biến màu sắc cho Light Mode và Dark Mode */
        :root {
            --bg-color: #f3f4f6;
            --card-bg: #ffffff;
            --text-main: #1f2937;
            --text-muted: #6b7280;
            --border-color: #e5e7eb;
            --th-bg: #f8fafc;
            --tr-hover: #f8fafc;
            --mac-bg: #eff6ff;
            --mac-color: #1d4ed8;
            --mac-border: #bfdbfe;
            --badge-bg: linear-gradient(135deg, #10b981 0%, #059669 100%);
            --alert-bg: #fef2f2;
            --alert-color: #991b1b;
            --alert-border: #fecaca;
        }

        [data-theme="dark"] {
            --bg-color: #0f172a;
            --card-bg: #1e293b;
            --text-main: #f8fafc;
            --text-muted: #94a3b8;
            --border-color: #334155;
            --th-bg: #1e293b;
            --tr-hover: #273548;
            --mac-bg: #1e3a8a;
            --mac-color: #93c5fd;
            --mac-border: #3b82f6;
            --badge-bg: linear-gradient(135deg, #059669 0%, #047857 100%);
            --alert-bg: #450a0a;
            --alert-color: #fca5a5;
            --alert-border: #7f1d1d;
        }

        * { box-sizing: border-box; transition: background-color 0.3s ease, color 0.3s ease, border-color 0.3s ease; }
        body {
            font-family: 'Inter', sans-serif;
            background-color: var(--bg-color);
            color: var(--text-main);
            margin: 0;
            padding: 40px 20px;
        }
        .dashboard-container { max-width: 1200px; margin: 0 auto; }
        
        .navbar-card {
            background: var(--card-bg);
            padding: 20px 30px;
            border-radius: 16px;
            box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.05);
            display: flex;
            justify-content: space-between;
            align-items: center;
            margin-bottom: 25px;
            border: 1px solid var(--border-color);
        }
        .brand-section { display: flex; align-items: center; gap: 16px; }
        .brand-logo {
            width: 70px;
            height: 70px;
            border-radius: 50%;
            object-fit: cover;
            box-shadow: 0 2px 8px rgba(183, 28, 28, 0.2);
            border: 2px solid white;
        }
        .brand-info h1 { font-size: 20px; font-weight: 700; color: var(--text-main); margin: 0 0 4px 0; }
        .brand-info p { font-size: 13px; color: var(--text-muted); margin: 0; }
        
        .nav-actions { display: flex; align-items: center; gap: 15px; }
        
        .stats-badge {
            background: var(--badge-bg);
            color: white;
            padding: 10px 20px;
            border-radius: 30px;
            font-weight: 600;
            font-size: 14px;
            box-shadow: 0 4px 12px rgba(16, 185, 129, 0.25);
            display: flex;
            align-items: center;
            gap: 8px;
        }

        /* Nút chuyển đổi Theme sáng/tối */
        .theme-toggle-btn {
            background: var(--th-bg);
            border: 1px solid var(--border-color);
            color: var(--text-main);
            width: 45px;
            height: 45px;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            cursor: pointer;
            font-size: 18px;
            box-shadow: 0 2px 8px rgba(0,0,0,0.05);
        }
        .theme-toggle-btn:hover { transform: scale(1.05); }

        #alert-box {
            background: var(--alert-bg);
            color: var(--alert-color);
            border: 1px solid var(--alert-border);
            padding: 16px 20px;
            border-radius: 12px;
            margin-bottom: 20px;
            display: none;
            font-weight: 500;
            font-size: 14px;
        }
        .table-card {
            background: var(--card-bg);
            border-radius: 16px;
            box-shadow: 0 4px 20px -2px rgba(0, 0, 0, 0.05);
            overflow: hidden;
            border: 1px solid var(--border-color);
        }
        table { width: 100%; border-collapse: collapse; text-align: left; }
        th {
            background-color: var(--th-bg);
            color: var(--text-muted);
            font-weight: 600;
            font-size: 13px;
            text-transform: uppercase;
            padding: 16px 24px;
            border-bottom: 1px solid var(--border-color);
        }
        td { padding: 16px 24px; border-bottom: 1px solid var(--border-color); font-size: 14px; color: var(--text-main); }
        tr:hover { background-color: var(--tr-hover); }
        
        .mac-text {
            font-family: 'JetBrains Mono', monospace;
            font-weight: 700;
            color: var(--mac-color);
            background: var(--mac-bg);
            padding: 6px 10px;
            border-radius: 6px;
            border: 1px solid var(--mac-border);
            letter-spacing: 0.5px;
        }

        .status-pill {
            display: inline-flex;
            align-items: center;
            gap: 6px;
            padding: 6px 12px;
            border-radius: 20px;
            font-weight: 600;
            font-size: 12px;
        }
        .status-ok { background-color: rgba(16, 185, 129, 0.1); color: #10b981; border: 1px solid rgba(16, 185, 129, 0.2); }
        .status-pending { background-color: rgba(245, 158, 11, 0.1); color: #f59e0b; border: 1px solid rgba(245, 158, 11, 0.2); }
        .spinner {
            width: 12px;
            height: 12px;
            border: 2px solid rgba(245, 158, 11, 0.3);
            border-radius: 50%;
            border-top-color: #f59e0b;
            animation: spin 0.8s linear infinite;
        }
        @keyframes spin { to { transform: rotate(360deg); } }
        .empty-state { text-align: center; padding: 40px; color: var(--text-muted); font-style: italic; }
    </style>
</head>
<body>
    <div class="dashboard-container">
        <div class="navbar-card">
            <div class="brand-section">
                <img src="data:image/png;base64,__LOGO_PLACEHOLDER__" alt="CLG Logo" class="brand-logo">
                <div class="brand-info">
                    <h1>Qr Grandstream by CLG</h1>
                    <p>Qr Grandstream  for Inventory Management</p>
                </div>
            </div>
            <div class="nav-actions">
                <div class="stats-badge">
                    <span>Tổng số thiết bị:</span>
                    <span id="total-count" style="font-size: 16px;">0</span>
                </div>
                <!-- Nút chuyển theme -->
                <button class="theme-toggle-btn" id="themeToggleBtn" onclick="toggleTheme()" title="Chuyển đổi giao diện Sáng/Tối">🌙</button>
            </div>
        </div>

        <div id="alert-box"></div>
        
        <div class="table-card">
            <table>
                <thead>
                    <tr>
                        <th>STT</th>
                        <th>MAC Address</th>
                        <th>Serial Number</th>
                        <th>Password</th>
                        <th>Thời Gian Quét</th>
                        <th>Trạng Thái</th>
                    </tr>
                </thead>
                <tbody id="table-body">
                    <tr><td colspan="6" class="empty-state">Đang kết nối cơ sở dữ liệu và chờ tín hiệu quét...</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <script>
        // Xử lý chuyển đổi giao diện Sáng / Tối có lưu LocalStorage
        function setTheme(theme) {
            document.documentElement.setAttribute('data-theme', theme);
            localStorage.setItem('theme', theme);
            let btn = document.getElementById('themeToggleBtn');
            btn.innerHTML = theme === 'dark' ? '☀️' : '🌙';
        }

        function toggleTheme() {
            let currentTheme = document.documentElement.getAttribute('data-theme');
            setTheme(currentTheme === 'dark' ? 'light' : 'dark');
        }

        // Khởi tạo theme khi load trang
        (function() {
            let savedTheme = localStorage.getItem('theme') || 'light';
            setTheme(savedTheme);
        })();

        function loadData() {
            fetch('/api/view_data')
                .then(res => res.json())
                .then(data => {
                    document.getElementById('total-count').innerText = data.total;
                    
                    let alertBox = document.getElementById('alert-box');
                    if (data.duplicates && data.duplicates.length > 0) {
                        let latestDup = data.duplicates[data.duplicates.length - 1];
                        alertBox.innerHTML = `⚠️ <strong>CẢNH BÁO TRÙNG LẶP:</strong> Thiết bị có địa chỉ MAC <span class="mac-text">${latestDup.mac}</span> đã tồn tại trong hệ thống! Thời gian quét: ${latestDup.time}`;
                        alertBox.style.display = "block";
                    } else {
                        alertBox.style.display = "none";
                    }

                    let tbody = document.getElementById('table-body');
                    tbody.innerHTML = "";
                    
                    if (data.items.length === 0) {
                        tbody.innerHTML = '<tr><td colspan="6" class="empty-state">Chưa có dữ liệu thiết bị nào được ghi nhận.</td></tr>';
                        return;
                    }

                    data.items.forEach(row => {
                        let tr = document.createElement('tr');
                        let statusHtml = (row[5] === 'Đã lưu') 
                            ? '<span class="status-pill status-ok">✓ Đã lưu Excel</span>' 
                            : '<span class="status-pill status-pending"><div class="spinner"></div> Đang đồng bộ</span>';
                        
                        tr.innerHTML = `
                            <td>${row[0]}</td>
                            <td><span class="mac-text">${row[1]}</span></td>
                            <td>${row[2]}</td>
                            <td>${row[3]}</td>
                            <td style="color: var(--text-muted); font-size: 13px;">${row[4]}</td>
                            <td>${statusHtml}</td>
                        `;
                        tbody.appendChild(tr);
                    });
                })
                .catch(err => console.error("Lỗi đồng bộ dữ liệu: ", err));
        }
        
        loadData();
        setInterval(loadData, 2000);
    </script>
</body>
</html>
"""

@app.route('/')
def index():
    IP = get_local_ip()
    return f"""
    <div style="font-family: 'Inter', sans-serif; text-align: center; margin-top: 80px; color: #1f2937;">
        <h2 style="color: #b71c1c;">Hệ thống Qr Grandstream by CLG đang hoạt động!</h2>
        <p style="color: #6b7280;">Webhook URL cấu hình trên ứng dụng barKoder:</p>
        <h3 style="background: #f3f4f6; display: inline-block; padding: 10px 20px; border-radius: 8px; color: #2563eb;">http://{IP}:5000/api/submit</h3>
        <div style="margin-top: 30px;">
            <a href="/view" style="background: #10b981; color: white; padding: 12px 24px; border-radius: 8px; text-decoration: none; font-weight: 600; box-shadow: 0 4px 12px rgba(16,185,129,0.3);">Mở Bảng Dashboard Live View</a>
        </div>
    </div>
    """

@app.route('/view')
def view_data():
    rendered_html = VIEW_TEMPLATE.replace('__LOGO_PLACEHOLDER__', LOGO_BASE64)
    return render_template_string(rendered_html)

@app.route('/api/view_data')
def api_view_data():
    init_excel_file()
    data_list = []
    with excel_lock:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]: 
                    data_list.append(list(row) + ["Đã lưu"])
        except Exception:
            pass
            
    global pending_queue, recent_duplicates
    current_stt = len(data_list) + 1
    for item in pending_queue:
        data_list.append([current_stt, item[0], item[1], item[2], item[3], "Chờ đồng bộ"])
        current_stt += 1
    data_list.reverse() 

    return jsonify({
        "total": len(data_list),
        "pending_count": len(pending_queue),
        "items": data_list,
        "duplicates": recent_duplicates
    })

@app.route('/api/submit', methods=['POST'])
def submit():
    global pending_queue, recent_duplicates
    init_excel_file()
    try:
        req_data = request.json or request.form or request.values
        raw_data = req_data.get('data', [])
        if isinstance(raw_data, str):
            try:
                raw_data = json.loads(raw_data)
            except:
                pass

        codes_batch = []
        if isinstance(raw_data, list):
            for item in raw_data:
                if isinstance(item, dict):
                    val = item.get('value', '').strip()
                    is_b64 = item.get('base64', False)
                    if val:
                        if is_b64:
                            try:
                                decoded_bytes = base64.b64decode(val)
                                val = decoded_bytes.decode('utf-8', errors='ignore').strip()
                            except:
                                pass
                        if val:
                            codes_batch.append(val)

        if len(codes_batch) < 3:
            return jsonify({
                "status": True, 
                "message": f"Đã nhận {len(codes_batch)} mã. Cần quét đủ 3 mã!"
            })

        target_codes = codes_batch[:3]
        mac, sn, password = "", "", ""
        for code in target_codes:
            cleaned = code.replace(":", "").replace("-", "").strip()
            if re.match(r'^[0-9A-Fa-f]{12}$', cleaned):
                mac = cleaned
            elif cleaned.isdigit() or len(cleaned) == 10:
                sn = cleaned
            else:
                password = cleaned

        unassigned = [c for c in target_codes if c not in [mac, sn, password]]
        if not mac and unassigned: mac = unassigned.pop(0)
        if not sn and unassigned: sn = unassigned.pop(0)
        if not password and unassigned: password = unassigned.pop(0)

        mac = mac.upper()
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with excel_lock:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
            
            for row in range(2, ws.max_row + 1):
                cell_val = ws.cell(row=row, column=2).value
                if cell_val and str(cell_val).strip().upper() == mac:
                    recent_duplicates = [{"mac": mac, "time": current_time}]
                    print(f"[!] PHÁT HIỆN TRÙNG MAC: {mac}")
                    return jsonify({
                        "status": False, 
                        "message": f"TRÙNG MAC: {mac} đã có trong hệ thống!"
                    })
            
            recent_duplicates.clear()

            ws.append([ws.max_row, mac, sn, password, current_time])
            wb.save(EXCEL_FILE)
            total_dev = ws.max_row - 1

        print(f"[✓] Đã lưu thành công! Tổng số: {total_dev}")
        return jsonify({
            "status": True, 
            "message": f"ĐÃ LƯU THÀNH CÔNG! (Tổng: {total_dev})"
        })

    except Exception as e:
        print(f"[x] Lỗi xử lý webhook: {str(e)}")
        return jsonify({"status": False, "message": f"Lỗi server: {str(e)}"})

def get_local_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(('8.8.8.8', 1))
        ip = s.getsockname()[0]
    except:
        ip = '127.0.0.1'
    finally:
        s.close()
    return ip

if __name__ == '__main__':
    init_excel_file()
    IP = get_local_ip()
    print("\n" + "="*70)
    print(" 🚀 HỆ THỐNG Qr Grandstream by CLG ĐÃ CÓ DARK/LIGHT MODE ")
    print("="*70)
    print(f"👉   http://{IP}:5000/api/submit")
    print("="*70)
    
    app.run(host='0.0.0.0', port=5000, threaded=True)